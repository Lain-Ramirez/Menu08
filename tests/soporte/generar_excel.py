"""Diligencia las cuatro hojas de la plantilla y añade resumen, evidencias y fuentes."""
from pathlib import Path
from collections import Counter
from copy import copy
import json,math,textwrap,hashlib,subprocess,os
import openpyxl
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.drawing.image import Image
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.chart import BarChart,Reference
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as col

ROOT=Path(__file__).resolve().parent.parent;EV=ROOT/'evidencias/actual';REPO=ROOT.parent
DOC=Path(os.environ['MENU08_ENTREGA_DIR']).resolve()
EVIDENCE_REF=os.environ.get('MENU08_EVIDENCE_REF','Lain')
S=json.loads((EV/'resumen.json').read_text());R=json.loads((EV/'ejecuciones.json').read_text());ACTUALES=len(R)
INICIAL=json.loads((ROOT/'evidencias/inicial/ejecuciones.json').read_text())
for id in ['CP-CARTA-08','CP-SEG-15','CP-AUTO-PM-01','CP-AUTO-PM-03']:
    old=next(x.copy() for x in INICIAL if x['caso']==id and x['estado']=='Fallida')
    old['ejecucion']='EX-BASE-'+str(len(R)-ACTUALES+1).zfill(3);old['historica']=True;old['caso_historico']=json.loads((ROOT/'evidencias/inicial'/(id+'.json')).read_text())['caso'];R.append(old)
CASES=json.loads((ROOT/'soporte/casos_entrega.json').read_text());BY={x['id']:x for x in CASES};FINAL={x['caso']:x for x in S['resultados']}
P='AA2502';G='725304';BG='FFF8F4';ALT='F5EBE4';INK='221A15';LINE='D7C2B7';GOLD='F5DF93';RED='FFDAD6';ORANGE='FFDCC2';WHITE='FFFFFF'
STATE={'Aprobada':GOLD,'Fallida':RED,'Pendiente de seguimiento':ORANGE}
W=openpyxl.load_workbook(DOC/'GA9_AA3_EV01_PLANTILLA_APRENDICES.xlsx')
W.worksheets[0].title='Diseño de Casos de Pruebas';W.worksheets[1].title='Ambiente de Pruebas'
for sh in W:
    for merged in list(sh.merged_cells.ranges):sh.unmerge_cells(str(merged))
    sh.delete_rows(6,sh.max_row)
    sh.data_validations.dataValidation=[]
    sh.sheet_view.showGridLines=False;sh.sheet_properties.pageSetUpPr.fitToPage=True
    sh.sheet_properties.outlinePr.summaryRight=False
    sh.sheet_properties.tabColor=P;sh.freeze_panes='D6';sh.sheet_view.zoomScale=70

def cell(sh,row,c,value,fill=None,bold=False,size=11,color=INK):
    z=sh.cell(row,c,value);z.font=Font(name='Aptos',size=size,bold=bold,color=color);z.alignment=Alignment(vertical='top',wrap_text=True);z.fill=PatternFill('solid',fgColor=fill or BG);z.border=Border(bottom=Side(style='hair',color=LINE));return z
def bar(sh,row,start,end,text,fill=P,size=14):
    sh.merge_cells(start_row=row,start_column=start,end_row=row,end_column=end)
    cell(sh,row,start,text,fill,True,size,WHITE if fill in [P,G] else INK);sh.row_dimensions[row].height=29
def hyperlink(c,target):c.hyperlink=target;c.font=Font(name='Aptos',size=c.font.sz or 11,color=P,underline='single',bold=c.font.bold)
def title(sh,n,subtitle):
    bar(sh,1,1,n,'MENU08 · GA9-220501096-AA3-EV01',P,18)
    bar(sh,2,1,n,subtitle,G,12)
    bar(sh,3,1,n,'Proyecto: Menu08 · CARTA, CAJA y SVP · commit '+S['commit'][:12]+' · ejecución automatizada 15/10/2026 UTC',ALT,11)
def headers(sh,row,labels,widths):
    for i,(name,width) in enumerate(zip(labels,widths),1):
        cell(sh,row,i,name,G,True,11,WHITE);sh.column_dimensions[col(i)].width=width
    sh.row_dimensions[row].height=37
def fitrow(sh,row,values,widths,maxheight=409):
    lines=max(sum(max(1,math.ceil(len(line)/(max(8,w-2)*1.05))) for line in str(v or '').split('\n')) for v,w in zip(values,widths))
    sh.row_dimensions[row].height=min(maxheight,max(54,lines*13+12))
def table(sh,labels,widths,rows):
    headers(sh,5,labels,widths)
    for i,vals in enumerate(rows,6):
        for j,v in enumerate(vals,1):cell(sh,i,j,v,BG if i%2==0 else ALT,size=10)
        fitrow(sh,i,vals,widths)
    sh.auto_filter.ref=f'A5:{col(len(labels))}{5+len(rows)}';sh.print_title_rows='1:5';sh.print_options.horizontalCentered=True
    sh.page_setup.orientation='landscape';sh.page_setup.paperSize=sh.PAPERSIZE_A3;sh.page_setup.fitToWidth=1;sh.page_setup.fitToHeight=0
    sh.print_area=f'A1:{col(len(labels))}{5+len(rows)}';sh.oddFooter.center.text='Menu08 · AA3 · Página &P de &N'

RESP=lambda cid:'Jovanny Medina Cifuentes (presentación y revisión)' if cid.startswith(('CP-SEG','CP-AUTO-PM')) or cid=='CP-CARTA-31' else 'Lain Sthid Ramirez Rueda (presentación y revisión)'
sh=W.worksheets[0];title(sh,13,'Diseño de casos de pruebas · trazabilidad con el plan AA2 y complementos AA3');bar(sh,4,1,13,'96 casos heredados del plan AA2 + 5 complementarios de automatización. Estado actual: hoja Resumen.',GOLD,11)
labels=[sh.cell(5,c).value for c in range(1,14)]
widths=[19,17,42,37,42,25,48,47,58,64,32,28,13]
rows=[]
for c in CASES:
    # Los antecedentes son una planificación; el estado observado está en Registro.
    rows.append([c[k] for k in ['id','modulo','requisito','nombre','objetivo','tipo','precondiciones','datos','pasos','esperado','herramienta']]+[RESP(c['id']),c['prioridad']])
table(sh,labels,widths,rows)
for i,c in enumerate(CASES,6):
    sh.cell(i,3).comment=Comment('Trazabilidad heredada de AA2; contrastada con el repositorio fijado al commit '+S['commit']+'.\n'+str(c.get('fuentes','')),'Menu08 QA')
    hyperlink(sh.cell(i,1),"#'Evidencias'!A"+str(116+(i-6)*26))

INC={
 'INC-01':dict(cases=['CP-CARTA-08','CP-AUTO-PM-03'],severity='Media',state='Verificada local',description='Producto no disponible visible en la carta. Reproduce DEF-01 del repositorio.',expected='Ocultarlo según criterio AA2 / aserción de la colección original.',actual='La carta lo muestra atenuado y con «No disponible».',steps='Ingresar food_truck; desactivar un producto; abrir /carta/truck-de-pruebas. La colección pública también rechaza Entrada agotada.',owner='Equipo Menu08 · catálogo',action='Corrección realizada: catalogoCarta filtra disponible=1. Reejecución HTTP y DOM en ambos navegadores: producto oculto al público y conservado en el panel. Carpeta pública Newman: 19 aserciones sin fallos.',version=S['commit'],retest='CP-CARTA-08'),
 'INC-02':dict(cases=['CP-AUTO-PM-01'],severity='Media',state='Verificada local',description='El recorrido Postman original usa IDs literales y cambia de rol sin cerrar sesión.',expected='20 peticiones originales con sus 36 aserciones sin fallos.',actual='20 peticiones realizadas; 15 aserciones fallidas. La venta usa IDs que no corresponden al banco y GET /ingresar redirige con sesión activa.',steps='Restaurar esquema, datos iniciales, Festín y pruebas. Ejecutar carpeta 1 con Newman --ignore-redirects. Revisar venta y cambios de rol.',owner='Equipo Menu08 · colección de pruebas',action='Corrección en la colección principal y en la muestra Menu08_GA9_VIDEO.postman_collection.json: descubrir IDs del HTML y añadir GET /salir antes de cada cambio de rol. Reejecución: 22 peticiones, 40 aserciones, 0 fallos.',version=S['commit'],retest='CP-AUTO-PM-02'),
 'INC-03':dict(cases=['CP-SEG-15'],severity='Alta',state='Pendiente de despliegue',description='El hosting permite HTTP y no entrega la redirección 301 a HTTPS requerida.',expected='HTTP 301 y Location=https://adso.menu08.com/carta/festin-rodante.',actual='HTTP 200, sin cabecera Location. HTTPS responde y valida TLS, pero no fuerza la migración desde HTTP.',steps='GET http://adso.menu08.com/carta/festin-rodante con redirecciones desactivadas; comparar código y Location.',owner='Equipo Menu08 · responsable de despliegue por asignar',action='Corrección realizada en ambas reglas .htaccess y antes de iniciar sesión en PHP. Apache HTTP/TLS y el front controller reales verifican 301, ruta, consulta y ausencia de bucles. Falta desplegar el PR y comprobar el hosting; no se afirma despliegue.',version=S['commit']+' (rama; hosting sin acreditar)',retest='CP-SEG-15'),
 'INC-04':dict(cases=['CP-AUTO-PM-03'],severity='Baja',state='Verificada local',description='La aserción de recursos remotos en /componentes interpreta cualquier URL absoluta como externa.',expected='Detectar recursos de un origen diferente; admitir recursos del propio origen.',actual='La aserción rechaza enlaces absolutos generados con url_base aunque apuntan al mismo servidor. La carpeta pública tiene 2 fallos: este y el de INC-01.',steps='Ejecutar carpeta 2. Abrir el test «No enlaza ningun recurso remoto» y comparar el origen de las URL de recursos con baseUrl.',owner='Equipo Menu08 · colección de pruebas',action='Corrección realizada en la colección principal: compara el origen de script, link e img con baseUrl. Permite URL absolutas del mismo servidor. Reejecución de la carpeta pública: 11 peticiones, 19 aserciones, 0 fallos.',version=S['commit'],retest='CP-AUTO-PM-03'),
 'SEG-01':dict(cases=['CP-CAJA-12'],severity='Baja',state='Pendiente de seguimiento',description='Falta exportación PDF desde el diálogo gráfico de Firefox.',expected='Comprobante imprimible revisado en ambos navegadores.',actual='Estilos print comprobados en ambos motores. PDF de 80 mm generado con Chromium; Firefox headless no exporta por el diálogo nativo.',steps='Abrir el comprobante en Firefox de escritorio, aceptar la vista previa, seleccionar PDF y guardar.',owner='Lain Sthid Ramirez Rueda',action='En la estación de grabación, exportar el PDF en Firefox y verificar consecutivo, ítems y total. Adjuntar archivo y cerrar seguimiento.',version='No aplica: comprobación física/gráfica',retest=''),
 'SEG-02':dict(cases=['CP-SVP-19'],severity='Baja',state='Pendiente de seguimiento',description='Lectura a 3 metros requiere monitor físico de 24 pulgadas.',expected='Números públicos legibles a 3 metros; controles internos >=32 px y >=56 px.',actual='Resoluciones, temas y medidas internas comprobados. No se dispone de verificación de distancia física.',steps='Abrir /turnos/truck-de-pruebas a pantalla completa en monitor de 24 pulgadas; observar los números a 3 metros.',owner='Jovanny Medina Cifuentes y Lain Sthid Ramirez Rueda',action='Grabar la prueba física y registrar aceptación antes de aprobar por completo CP-SVP-19.',version='No aplica: prueba física',retest=''),
 'SEG-03':dict(cases=['CP-SVP-20'],severity='Media',state='Pendiente de seguimiento',description='Aceptación de negocio de cerrar un turno con órdenes pendientes sin resolver en el repositorio.',expected='Registrar comportamiento real y conservar la decisión pendiente.',actual='Al cerrar, turno=null y ordenes=[] en SVP; la orden persiste en la BD.',steps='Crear una orden pendiente; cerrar turno; consultar /svp/ordenes y la fila de ordenes.',owner='Equipo Menu08 · decisión de producto por asignar',action='Definir si se debe impedir el cierre o mostrar pendientes históricos. Después ajustar criterio y reejecutar; no se afirma aprobación de negocio.',version='No aplicada',retest=''),
}
def incs(cid):return ', '.join(k for k,v in INC.items() if cid in v['cases'])
def observed(r):
    good=sum(x['cumple'] for x in r['checks']);bad=[x for x in r['checks'] if not x['cumple']]
    text=f"{good}/{len(r['checks'])} comparaciones cumplen."
    if bad:text+='\n'+'\n'.join(x['comprobacion']+': esperado '+str(x['esperado'])+'; obtenido '+str(x['obtenido']) for x in bad)
    else:text+='\n'+'\n'.join(x['comprobacion']+': '+str(x['obtenido']) for x in r['checks'][:8])
    if len(r['checks'])>8:text+='\nDetalle completo en el informe y JSON vinculados.'
    if r['pendientes']:text+='\nPENDIENTE: '+' / '.join(r['pendientes'])
    return text
sh=W['Registro de ejecución'];title(sh,13,'Registro de ejecución · esperado frente a observado · cada fila corresponde a una ejecución real');bar(sh,4,1,13,'Ejecutor: automatización local. Responsables indicados para revisión y presentación. Fechas UTC; hosting solo lectura.',GOLD,11)
labels=[sh.cell(5,c).value for c in range(1,14)];widths=[17,20,25,25,29,49,64,78,25,45,24,34,65]
rows=[]
for r in R:
    c=r.get('caso_historico',BY[r['caso']]);details='\n'.join(r['notas'])
    if len(details)>1500:details=details[:1500]+'\nDetalle íntegro en JSON.'
    rows.append([r['ejecucion'],r['caso'],r['fecha'][:19].replace('T',' ')+' UTC',r['commit'],'Automatizada · revisión: '+RESP(r['caso']),c['datos'],c['esperado'],observed(r),r['estado'],('evidencias/inicial/' if r.get('historica') else 'evidencias/actual/')+r['caso']+'.json',incs(r['caso']),r.get('herramienta','Automatización'),details])
table(sh,labels,widths,rows)
for i,r in enumerate(R,6):
    hyperlink(sh.cell(i,10),('evidencias/inicial/' if r.get('historica') else 'evidencias/actual/')+r['caso']+'.json');cell(sh,i,9,r['estado'],STATE[r['estado']],True,10)
dv=DataValidation(type='list',formula1='"Aprobada,Fallida,Pendiente de seguimiento"');sh.add_data_validation(dv);dv.add(f'I6:I{5+len(R)}')

sh=W['Incidencias y seguimiento'];title(sh,14,'Incidencias, acciones correctivas y reejecuciones verificables');bar(sh,4,1,14,'INC = defectos reproducidos; SEG = validación o decisión pendiente. La corrección local de Postman está comprobada.',GOLD,11)
labels=[sh.cell(5,c).value for c in range(1,15)];widths=[17,24,25,54,62,55,67,15,28,37,85,40,26,47];rows=[]
for key,v in INC.items():
    origin=[r['ejecucion'] for r in R if r['caso'] in v['cases'] and r.get('historica')]
    re=[r['ejecucion'] for r in R if r['caso']==v['retest'] and not r.get('historica')]
    rows.append([key,', '.join(origin),', '.join(v['cases']),v['description'],v['steps'],v['expected'],v['actual'],v['severity'],v['state'],v['owner'],v['action'],v['version'],', '.join(re) if re else 'Pendiente; no se declara verificación','evidencias/actual/'+v['cases'][0]+'.html'])
table(sh,labels,widths,rows)
for i in range(6,6+len(rows)):hyperlink(sh.cell(i,14),rows[i-6][13])
for column,options in [('H','Alta,Media,Baja'),('I','Abierta,Verificada local,Pendiente de despliegue,Pendiente de seguimiento')]:
    dv=DataValidation(type='list',formula1='"'+options+'"');sh.add_data_validation(dv);dv.add(f'{column}6:{column}{5+len(rows)}')

# Ambiente: conservar los cuatro bloques y sus posiciones de la plantilla.
sh=W['Ambiente de Pruebas'];sh.delete_rows(1,sh.max_row);sh.freeze_panes='B5'
for i,w in enumerate([32,69,47,94],1):sh.column_dimensions[col(i)].width=w
bar(sh,1,1,4,'MENU08 · DEFINICIÓN Y CONFIGURACIÓN DEL AMBIENTE DE PRUEBAS',P,17)
bar(sh,2,1,4,'1. Componentes Técnicos del Entorno',G,13)
headers(sh,4,['Componente Técnico','Descripción / Especificación','Versión Seleccionada','Observaciones / Propósito'],[32,69,47,94])
env=json.loads((EV/'entorno.json').read_text());nav=json.loads((EV/'navegadores.json').read_text())
envrows=[['Sistema Operativo',env['so'],'Linux / Ubuntu 24.04.5 LTS','Servidor y ejecutor de pruebas. No se presenta como Windows ni Android físico.'],['Frontend','HTML, CSS y JavaScript propios del repositorio','Tokens MD3 Brasa; sin framework de frontend','Rojo #AA2502, dorado #725304 y crema #FFF8F4, tomados de ADSO.menu08.com/recursos/css/md3.css.'],['Backend',env['php'],'PHP 8.3.6','PDO MySQL; front controller y MVC del commit verificado.'],['Sistema Gestor BD','Instancia MySQL aislada en Docker',env['mysql'],'127.0.0.1:13306; base menu08 local, 9 tablas InnoDB; utf8mb4_unicode_ci.'],['Servidor Web / Ejecución','Servidor integrado de PHP','127.0.0.1:8088','Raíz pública ADSO.menu08.com; configuración y SQL fuera de esa raíz. Este ciclo local no prueba LiteSpeed.'],['Navegadores de Prueba','Motores reales automatizados con Playwright','Chromium '+nav['chromium']+' / Firefox '+nav['firefox'],'Anchos 320, 360, 768 y 1280; temas claro y oscuro. Sin validación de monitor físico ni Android.']]
for i,row in enumerate(envrows,5):
    for j,v in enumerate(row,1):cell(sh,i,j,v,ALT if i%2 else BG,size=11)
    fitrow(sh,i,row,[32,69,47,94])
bar(sh,12,1,4,'2. Herramientas de Pruebas Utilizadas',G,13)
headers(sh,14,['Herramienta','Tipo de Prueba','Aplicación / Alcance','Registro verificable'],[32,69,47,94])
tools=[['Ejecutor PHP del repositorio','Unitaria','62 comprobaciones; 0 fallos','evidencias/unitarias.txt'],['Postman / Newman 6.2.1','Servicios / integración','Colección corregida: carpetas 1, 2 y 8; muestra del video','22 + 11 + 8 peticiones de la colección; 22 de la muestra. Fallos previos conservados en evidencias/inicial.'],['Playwright '+env['playwright'],'E2E / presentación / impresión','Chromium y Firefox','Capturas reales, comparación DOM y flujo completo por interfaz.'],['Python requests / PHP PDO / cliente mysql','HTTP, integridad y seguridad','Comparar HTTP y estado de las tablas','Respuestas y checks en evidencias/CP-*.json. Fallo SQL inducido retirado.']]
for i,row in enumerate(tools,15):
    for j,v in enumerate(row,1):cell(sh,i,j,v,ALT,size=11)
    fitrow(sh,i,row,[32,69,47,94])
bar(sh,19,1,4,'3. Configuraciones y Variables de Entorno (.env.testing)',G,13)
bar(sh,20,1,4,'Menu08 usa configuracion/configuracion.php, derivada de configuracion.ejemplo.php. No carga .env.testing.',GOLD,11)
configs={21:'Configuración REAL del ciclo local; no contiene credenciales del hosting.',22:"entorno = produccion (errores internos en bitácora)",23:"url_base = http://127.0.0.1:8088; zona_horaria = America/Bogota",24:"sesion.nombre = menu08_sesion; vida_minutos = 120; solo_https = false (HTTP local)",25:'Sesiones locales: soporte/sesiones; subidas: ADSO.menu08.com/subidas',26:'Configuración de la base local de pruebas',27:'base_datos.servidor = 127.0.0.1; puerto = 13306',28:'base_datos.nombre = menu08 (contenedor aislado)',29:'base_datos.usuario = root; contraseña local de demostración definida por soporte/configurar_local.py',30:'base_datos.cotejamiento = utf8mb4_unicode_ci; SET NAMES del cliente con ese mismo cotejamiento',31:'PDO: ERRMODE_EXCEPTION; ATTR_EMULATE_PREPARES=false; sesión MySQL time_zone=-05:00',32:'Hosting: solo consultas de lectura a HTTPS/HTTP. Sin acceso a bitácora ni verificación del commit desplegado.'}
for row,text in configs.items():bar(sh,row,1,4,text,ALT,11);sh.row_dimensions[row].height=32
bar(sh,34,1,4,'4. Datos de Prueba y Carga Inicial (Seeds / Mocks)',G,13)
headers(sh,36,['Recurso','Archivo / Script','Descripción de los Datos Cargados','Observaciones'],[32,69,47,94])
seeds=[['Estructura','menu08_app/basedatos/esquema.sql','9 tablas InnoDB','Se importa únicamente en la base del contenedor local.'],['Demostración','datos_iniciales.sql + datos_festin_rodante.sql','Catálogo Festín Rodante de referencia','Se contrasta aislamiento desde la instancia local. No se alteró el hosting.'],['Banco de prueba','menu08_app/basedatos/datos_pruebas.sql','1 truck, 3 usuarios, 5 categorías, 11 productos, 4 paradas','Incluye producto agotado, gratuito, centavos y parada nocturna. Se restauró entre ciclos. IDs consultados por nombre/slug.'],['Instrumentación','soporte/ejecutar_*.py y soporte/puente.php','Comprobaciones derivadas del comportamiento del repositorio','No cambia archivos versionados del producto; prepara y contrasta la base local. Los responsables humanos revisan/presentan la evidencia.'],['Datos controlados','Antigüedad 9/10 minutos; trigger de fallo SQL; fecha de vencimiento de token local','Solo en la instancia aislada','Se registra la manipulación de preparación. Trigger retirado en finally. No se interpreta como un defecto real el fallo inducido.']]
for i,row in enumerate(seeds,37):
    for j,v in enumerate(row,1):cell(sh,i,j,v,ALT if i%2 else BG,size=11)
    fitrow(sh,i,row,[32,69,47,94])
sh.print_area='A1:D41';sh.page_setup.orientation='landscape';sh.page_setup.paperSize=sh.PAPERSIZE_A3;sh.page_setup.fitToWidth=1;sh.page_setup.fitToHeight=0

sh=W.create_sheet('Resumen');sh.sheet_view.showGridLines=False;sh.sheet_properties.tabColor=G
for c in range(1,14):sh.column_dimensions[col(c)].width=14
title(sh,13,'Documenta pruebas de software de acuerdo con la planificación · entrega AA3 EV01')
bar(sh,5,1,13,'Lain Sthid Ramirez Rueda · Jovanny Medina Cifuentes · Ficha 3235887',ALT,12)
bar(sh,6,1,13,'Fecha de ejecución: 15/10/2026 UTC · Fuente única del producto: github.com/Lain-Ramirez/Menu08',ALT,11)
bar(sh,8,1,3,'ENLACE AL VIDEO',G,12);sh.merge_cells('D8:M9');cell(sh,8,4,'',GOLD);sh['D8'].comment=Comment('Pegar aquí el enlace al video cuando esté grabado y publicado. Este espacio se deja intencionalmente vacío.','Menu08 QA');sh.row_dimensions[9].height=25
bar(sh,11,1,13,'Resultados por caso · cada caso conserva las ejecuciones que lo sustentan',P,13)
for row,name in enumerate(['Aprobada','Fallida','Pendiente de seguimiento'],13):
    n=S['estados'].get(name,0);cell(sh,row,1,name,STATE[name],True,12);sh.merge_cells(start_row=row,start_column=1,end_row=row,end_column=4);cell(sh,row,5,n,STATE[name],True,16);sh.row_dimensions[row].height=29
bar(sh,17,1,13,f"Plan AA2: 96 casos · Complementarios: 5 · Ejecuciones actuales: {ACTUALES} · Histórico fallido: 4 · Comparaciones: {S['comprobaciones']}",ALT,11)
bar(sh,19,1,13,'Correcciones locales verificadas · Cierre global pendiente de despliegue, pruebas físicas y decisión de negocio.',ORANGE,13);sh.row_dimensions[19].height=32
bar(sh,21,1,13,'Situación real que debe explicarse en el video',G,12)
notes=[
 'Los cuatro casos antes fallidos tienen correcciones y reejecuciones locales. El historial conserva los resultados iniciales y su commit.',
 'CARTA: agotados ocultos al público y conservados en el panel. Postman: recorrido de 22 peticiones/40 aserciones; públicas 11/19; móvil 8/25, sin fallos.',
 'HTTPS: ambas reglas .htaccess y el front controller real pasan en Apache/PHP locales. El hosting todavía requiere desplegar el PR; su HTTP observado se registra aparte.',
 'Pendientes: despliegue HTTPS, PDF desde Firefox gráfico, lectura a 3 metros en monitor de 24 pulgadas y definición del cierre con órdenes pendientes.',
 'DEF-02 histórico (desborde a 320 px): no reproducido. Las capturas de aplicación y los informes automatizados están identificados y enlazados en GitHub.',
 'La carpeta documental contiene los Excel y el guion. Código de pruebas, registros y capturas completas viven en Menu08/tests; el libro incorpora miniaturas sin dependencias locales.',
]
for row,text in enumerate(notes,23):bar(sh,row,1,13,text,ALT,11);sh.row_dimensions[row].height=37
bar(sh,30,1,13,'Navegación de la entrega',G,12)
for row,name in enumerate(['Diseño de Casos de Pruebas','Ambiente de Pruebas','Registro de ejecución','Incidencias y seguimiento','Evidencias','Fuentes'],32):
    bar(sh,row,1,13,name,GOLD,12);hyperlink(sh.cell(row,1),"#'"+name+"'!A1")
bar(sh,39,1,13,'Guion: GUION_VIDEO_JOVANNY_LAIN.md · Pruebas y evidencias completas: repositorio Menu08, carpeta tests.',ALT,11)
sh.freeze_panes='A11';sh.sheet_view.zoomScale=90;sh.print_area='A1:M39';sh.page_setup.orientation='landscape';sh.page_setup.paperSize=sh.PAPERSIZE_A3;sh.page_setup.fitToWidth=1;sh.page_setup.fitToHeight=1

sh=W.create_sheet('Evidencias');sh.sheet_view.showGridLines=False;sh.sheet_properties.tabColor=P;sh.freeze_panes='C6'
for c in range(1,14):sh.column_dimensions[col(c)].width=13
title(sh,13,'Índice y capturas incrustadas · informe automatizado + captura original de aplicación cuando corresponde')
bar(sh,4,1,13,'Cada enlace abre el informe completo y su JSON. Las miniaturas permiten revisar sin archivos externos.',GOLD,11)
headers(sh,5,['Caso','Estado','Nombre','Ejecuciones','Informe','JSON','Capturas','', '', '', '', '', ''],[21,27,48,30,20,20,15,13,13,13,13,13,13])
for i,r in enumerate(S['resultados'],6):
    start=116+(i-6)*26
    vals=[r['caso'],r['estado'],r['nombre'],', '.join(r['ejecuciones']),'Abrir informe','Abrir JSON',len(r['capturas'])]
    for j,v in enumerate(vals,1):cell(sh,i,j,v,STATE[r['estado']] if j==2 else ALT,size=10)
    sh.row_dimensions[i].height=33;hyperlink(sh.cell(i,1),"#'Evidencias'!A"+str(start));hyperlink(sh.cell(i,5),'evidencias/actual/'+r['caso']+'.html');hyperlink(sh.cell(i,6),'evidencias/actual/'+r['caso']+'.json')
    bar(sh,start,1,13,r['caso']+' · '+r['nombre']+' · '+r['estado'],P,12)
    bar(sh,start+1,1,7,'INFORME AUTOMATIZADO DE LOS REGISTROS',GOLD,10);bar(sh,start+1,8,13,'CAPTURA ORIGINAL DE APLICACIÓN / HERRAMIENTA',GOLD,10)
    def pic(path,anchor,maxw,maxh):
        img=Image(path);ratio=min(maxw/img.width,maxh/img.height);img.width*=ratio;img.height*=ratio;sh.add_image(img,anchor)
    pic(EV/(r['caso']+'-informe.png'),'A'+str(start+2),900,390)
    if r['capturas']:
        selected=next((p for p in r['capturas'] if 'comprobante' in p or 'sondeo' in p or 'cierre' in p or 'producto-agotado' in p),r['capturas'][0]);pic(ROOT/selected,'H'+str(start+2),540,390)
    elif r['caso']=='CP-CARTA-13':pic(EV/'qr-local.png','H'+str(start+2),300,300)
    elif r['caso']=='CP-CARTA-31':pic(EV/'unitarias.png','H'+str(start+2),540,390)
    elif r['caso'].startswith('CP-AUTO-PM'):
        name={'CP-AUTO-PM-01':'postman-original-recorrido','CP-AUTO-PM-02':'postman-video-validado','CP-AUTO-PM-03':'postman-original-publicas','CP-AUTO-PM-04':'postman-original-movil'}[r['caso']]
        pic(EV/(name+'-final.png'),'H'+str(start+2),540,390)
    else:cell(sh,start+3,8,'Esta prueba se acredita con respuestas HTTP, consultas SQL o resultados CLI. Abrir el informe para consultar el registro completo.',ALT,size=11);sh.merge_cells(start_row=start+3,start_column=8,end_row=start+7,end_column=13)
    bar(sh,start+24,1,13,'Volver al índice de evidencias · '+str(len(r['checks']))+' comparaciones · '+', '.join(r['ejecuciones']),ALT,10);hyperlink(sh.cell(start+24,1),"#'Evidencias'!A5")
sh.auto_filter.ref='A5:G106';sh.sheet_view.zoomScale=60

sh=W.create_sheet('Fuentes');sh.sheet_view.showGridLines=False;sh.sheet_properties.tabColor=G;sh.freeze_panes='B6';title(sh,5,'Procedencia verificable · enlaces fijados al commit consultado y hashes de archivos originales')
headers(sh,5,['ID','Archivo / recurso del repositorio','Uso','Enlace inmutable','SHA-256'],[14,72,48,105,75])
sourcefiles=sorted(set(list((REPO/'docs').rglob('*.md'))+list((REPO/'menu08_app/aplicacion').rglob('*.php'))+list((REPO/'menu08_app/pruebas').rglob('*.php'))+list((REPO/'menu08_app/basedatos').glob('*.sql'))+list((REPO/'ADSO.menu08.com/recursos').rglob('*.css'))+list((REPO/'ADSO.menu08.com/recursos').rglob('*.js'))+[REPO/'POSTMAN.md',REPO/'README.md',REPO/'postman/Menu08.postman_collection.json',REPO/'menu08_app/configuracion/rutas.php']))
for i,path in enumerate(sourcefiles,6):
    rel=path.relative_to(REPO).as_posix();url='https://github.com/Lain-Ramirez/Menu08/blob/'+S['commit']+'/'+rel
    vals=[f'F-{i-5:03}',rel,'Fuente primaria del producto y su planificación; commit verificado',url,hashlib.sha256(path.read_bytes()).hexdigest()]
    for j,v in enumerate(vals,1):cell(sh,i,j,v,ALT if i%2 else BG,size=10)
    sh.row_dimensions[i].height=41;hyperlink(sh.cell(i,4),url)
sh.auto_filter.ref=f'A5:E{5+len(sourcefiles)}';sh.print_title_rows='1:5'
W.active=W.sheetnames.index('Resumen');W.calculation=openpyxl.workbook.properties.CalcProperties(calcId=191029,fullCalcOnLoad=True)
W.properties.creator='Lain Sthid Ramirez Rueda / Jovanny Medina Cifuentes';W.properties.title='Menu08 · GA9 AA3 EV01 · Ejecución real de pruebas';W.properties.subject='96 casos AA2, pruebas complementarias y evidencia verificable';W.properties.description='Plantilla AA3 diligenciada. Ejecución automatizada local y consultas de lectura del hosting. No se certifica aprobación global.'
for sheet in W:
    for row in sheet:
        for c in row:
            if c.hyperlink and not c.hyperlink.target.startswith(('#','http')):
                target=c.hyperlink.target
                if target.endswith('.html') and (ROOT/target.replace('.html','.json')).exists():
                    target=target.replace('.html','.json')
                c.hyperlink='https://github.com/Lain-Ramirez/Menu08/blob/'+EVIDENCE_REF+'/tests/'+target
out=DOC/'GA9-220501096-AA3-EV01_Menu08_Pruebas_Ejecutadas.xlsx';W.save(out)
print(out.name,out.stat().st_size,'bytes',len(R),'ejecuciones')
