"""Verifica el libro entregado y sus enlaces a los registros del repositorio."""
from pathlib import Path
import os,json,zipfile,re,hashlib
import openpyxl
ROOT=Path(__file__).resolve().parent.parent
DOC=Path(os.environ['MENU08_ENTREGA_DIR']).resolve()
EV=ROOT/'evidencias/actual'
W=openpyxl.load_workbook(DOC/'GA9-220501096-AA3-EV01_Menu08_Pruebas_Ejecutadas.xlsx')
S=json.loads((EV/'resumen.json').read_text())
T=openpyxl.load_workbook(DOC/'GA9_AA3_EV01_PLANTILLA_APRENDICES.xlsx')
assert len(S['resultados'])==101 and len({x['caso'] for x in S['resultados']})==101
assert W['Resumen']['D8'].value is None
assert len(W['Evidencias']._images)>=101
for old,new,n in [(T.worksheets[0],W.worksheets[0],13),(T.worksheets[2],W.worksheets[2],13),(T.worksheets[3],W.worksheets[3],14)]:
    assert [old.cell(5,i).value for i in range(1,n+1)]==[new.cell(5,i).value for i in range(1,n+1)]
links=0
for sheet in W:
    for row in sheet:
        for c in row:
            assert c.data_type!='e'
            if c.hyperlink:
                links+=1;t=c.hyperlink.target
                assert t.startswith(('#','https://github.com/Lain-Ramirez/Menu08/'))
                if '/tests/evidencias/' in t:
                    rel=t.split('/tests/',1)[1]
                    assert (ROOT/rel).is_file(),rel
assert zipfile.ZipFile(DOC/'GA9-220501096-AA3-EV01_Menu08_Pruebas_Ejecutadas.xlsx').testzip() is None
report={'resultado':'Validada','casos':101,'estados':S['estados'],'enlaces_excel':links,'imagenes_incrustadas':len(W['Evidencias']._images),'video':'Resumen!D8 vacío','limite':'Validación estructural, sin ejecutar Microsoft Excel/LibreOffice. El hosting y las pruebas físicas conservan sus seguimientos.'}
(EV/'validacion_entrega.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
manifest={f.name:hashlib.sha256(f.read_bytes()).hexdigest() for f in EV.iterdir() if f.is_file() and f.name!='SHA256.json'}
(EV/'SHA256.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
print(json.dumps(report,ensure_ascii=False,indent=2))
